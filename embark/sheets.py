"""Excel files the tool hands out: blank templates for the customer, prepared data for the consultant."""

from __future__ import annotations

import io
import zipfile

import frappe
from frappe import _
from frappe.utils import cstr
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from embark.engine import build_columns

INK = "171717"
MUTED = "E5E7EB"
TEMPLATE_ROWS = 5000


def template_xlsx(area, answers: dict[str, str] | None = None) -> bytes:
	"""A sheet with the step's headings, dropdowns for fixed choices and a guide tab."""
	columns = build_columns(area, answers)
	wb = Workbook()
	ws = wb.active
	ws.title = area.area_name[:31]

	for i, c in enumerate(columns, start=1):
		letter = get_column_letter(i)
		heading = c.label + (" *" if c.required else "")
		cell = ws.cell(row=1, column=i, value=heading)
		cell.font = Font(bold=True, color="FFFFFF" if c.required else INK)
		cell.fill = PatternFill("solid", fgColor=INK if c.required else MUTED)
		ws.column_dimensions[letter].width = max(16, len(heading) + 4)

		choices = c.choices or (["Yes", "No"] if c.fieldtype == "Check" else [])
		# Excel caps an inline list at 255 characters; longer lists stay free text.
		if choices and len(",".join(choices)) < 250:
			dv = DataValidation(type="list", formula1='"{}"'.format(",".join(choices)), allow_blank=True)
			dv.add(f"{letter}2:{letter}{TEMPLATE_ROWS}")
			ws.add_data_validation(dv)
	ws.freeze_panes = "A2"

	guide = wb.create_sheet(_("How to fill"))
	guide.append([_("Column"), _("Required"), _("What to enter"), _("Example")])
	for cell in guide[1]:
		cell.font = Font(bold=True)
	for c in columns:
		guide.append([c.label, _("Yes") if c.required else _("No"), describe(c), c.example])
	for letter, width in zip("ABCD", (24, 10, 70, 28), strict=True):
		guide.column_dimensions[letter].width = width
	for row in guide.iter_rows(min_row=2):
		row[2].alignment = Alignment(wrap_text=True, vertical="top")

	return _save(wb)


def typed_xlsx(area, rows: list[dict], answers: dict[str, str] | None = None) -> bytes:
	"""A sheet built from rows someone typed in, rather than a file they had.

	The rest of Embark only knows how to read a spreadsheet, so typing a few
	warehouses in makes one — and every check, fix and export then works on it
	exactly as it would on a customer's own file.
	"""
	columns = build_columns(area, answers or {})
	wb = Workbook()
	ws = wb.active
	ws.title = area.name[:31]
	ws.append([c.label for c in columns])
	for row in rows:
		ws.append([cstr(row.get(c.fieldname, "")).strip() for c in columns])
	return _save(wb)


def describe(c) -> str:
	if c.hint:
		return c.hint
	if c.fieldtype == "Check":
		return _("Yes or No.")
	if c.choices:
		return _("One of: {0}.").format(", ".join(c.choices))
	if c.fieldtype == "Link":
		return _("A {0}.").format(_(c.link_doctype))
	if c.fieldtype in ("Int", "Float"):
		return _("A number.")
	if c.fieldtype == "Date":
		return _("A date, such as 31-12-2026.")
	if c.fieldtype == "Email":
		return _("An email address.")
	if c.fieldtype == "Phone":
		return _("A phone number.")
	return ""


def prepared_zip(onboarding) -> bytes:
	"""Every checked sheet in ERPNext's shape, headed with the doctype's own labels.

	Headings match what ERPNext's Data Import expects, so these files load
	with no column mapping. Extra columns (contact, address) are kept at the
	end for the loader and are skipped by Data Import.
	"""
	buffer = io.BytesIO()
	with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
		zf.writestr("00 Company Details.xlsx", _company_sheet(onboarding))
		for n, row in enumerate(sorted(onboarding.areas, key=_sequence), start=1):
			if not row.upload:
				continue
			evaluated = frappe.get_doc("Embark Upload", row.upload).evaluate()
			zf.writestr(f"{n:02d} {row.data_area}.xlsx", _data_sheet(onboarding, evaluated))
	return buffer.getvalue()


def _sequence(row) -> int:
	return frappe.get_cached_value("Embark Data Area", row.data_area, "sequence") or 0


def _company_sheet(onboarding) -> bytes:
	wb = Workbook()
	ws = wb.active
	ws.title = _("Company")
	meta = frappe.get_meta("Embark Onboarding")
	for fieldname in (
		"company_name",
		"abbr",
		"country",
		"default_currency",
		"fiscal_year_start",
		"chart_of_accounts",
		"tax_id",
		"company_email",
		"company_phone",
		"company_address",
	):
		ws.append([_(meta.get_label(fieldname)), onboarding.get(fieldname)])
	ws.column_dimensions["A"].width = 28
	ws.column_dimensions["B"].width = 48
	return _save(wb)


def _data_sheet(onboarding, evaluated) -> bytes:
	meta = frappe.get_meta(evaluated.area.target_doctype)
	columns = evaluated.columns
	company_fields = [
		df.fieldname for df in meta.fields if df.fieldtype == "Link" and df.options == "Company" and df.reqd
	]

	wb = Workbook()
	ws = wb.active
	ws.title = evaluated.area.area_name[:31]
	ws.append(
		[_(meta.get_label(f)) for f in company_fields]
		+ [c.label if c.is_extra else _(meta.get_label(c.fieldname)) for c in columns]
	)
	for cell in ws[1]:
		cell.font = Font(bold=True)
	for r in evaluated.result.rows:
		ws.append([onboarding.company_name for _f in company_fields] + [r.get(c.fieldname) for c in columns])
	return _save(wb)


def _save(wb) -> bytes:
	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()
