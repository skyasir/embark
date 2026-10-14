"""The customer journey end to end, run as a real Website User.

bench --site <site> run-tests --module embark.tests.test_onboarding
"""

import io
import json
import typing
import zipfile
from unittest.mock import patch

import frappe
from openpyxl import Workbook, load_workbook

from embark import api, assistant, wizard
from embark.conditions import applies
from embark.engine import Column, build_columns, coerce, match_columns, split_header

try:  # v16
	from frappe.tests import IntegrationTestCase as TestCase
except ImportError:  # v15
	from frappe.tests.utils import FrappeTestCase as TestCase


def xlsx(rows) -> bytes:
	wb = Workbook()
	for r in rows:
		wb.active.append(r)
	out = io.BytesIO()
	wb.save(out)
	return out.getvalue()


class TestEngine(TestCase):
	def test_header_below_a_title_row(self):
		headers, rows = split_header(
			[["Customer List"], [], ["Name", "GSTIN"], ["Sunrise", "27AAA"], [None, None]]
		)
		self.assertEqual(headers, ["Name", "GSTIN"])
		self.assertEqual(rows, [(4, ["Sunrise", "27AAA"])])

	def test_columns_match_by_alias_and_fuzzy(self):
		cols = [
			Column("customer_name", "Customer Name", "Data", aliases=["Party Name"]),
			Column("tax_id", "Tax ID", "Data", aliases=["GSTIN"]),
			Column("customer_group", "Customer Group", "Link"),
		]
		self.assertEqual(
			match_columns(["Party Name", "GSTIN No.", "Customer Grp", "Notes"], cols),
			{"Party Name": "customer_name", "GSTIN No.": "tax_id", "Customer Grp": "customer_group"},
		)
		self.assertEqual(match_columns(["GSTIN", "Notes"], cols), {"GSTIN": "tax_id"})

	def test_coercion(self):
		check = Column("x", "X", "Check")
		self.assertEqual(coerce(check, "yes")[:2], (1, None))
		self.assertEqual(coerce(check, "maybe")[1], "BAD_YESNO")
		num = Column("x", "X", "Float")
		self.assertEqual(coerce(num, "₹1,250.50")[:2], (1250.5, None))
		phone = Column("x", "X", "Phone")
		self.assertEqual(coerce(phone, 9876543210.0)[:2], ("9876543210", None))
		choice = Column("x", "X", "Select", choices=["Company", "Individual"])
		self.assertEqual(coerce(choice, "individual")[:2], ("Individual", None))
		date = Column("x", "X", "Date")
		self.assertEqual(coerce(date, "03/04/2026")[0], "2026-04-03")
		required = Column("x", "X", "Data", required=True)
		self.assertEqual(coerce(required, "  ")[1], "REQUIRED_MISSING")
		defaulted = Column("x", "X", "Check", default="Yes")
		self.assertEqual(coerce(defaulted, None)[:2], (1, None))


class TestConditions(TestCase):
	def test_conditions(self):
		answers = {
			"keeps_stock": "yes",
			"business_type": "trading,manufacturing",
			"tracks_batches": "no",
			"unsure": "not_sure",
		}
		self.assertTrue(applies("", answers))
		self.assertTrue(applies("keeps_stock == yes", answers))
		self.assertFalse(applies("keeps_stock == no", answers))
		self.assertTrue(applies("business_type == manufacturing", answers))  # one of several
		self.assertTrue(applies("tracks_batches != yes", answers))
		self.assertTrue(applies("tracks_batches == yes or keeps_stock == yes", answers))
		self.assertFalse(applies("tracks_batches == yes and keeps_stock == yes", answers))
		# Unanswered, "not sure" and nonsense all leave the step visible.
		self.assertTrue(applies("never_asked == yes", answers))
		self.assertTrue(applies("unsure == yes", answers))
		self.assertTrue(applies("!!!", answers))


class TestOnboardingFlow(TestCase):
	def setUp(self):
		frappe.set_user("Administrator")
		frappe.db.delete("Embark Upload")
		frappe.db.delete("Embark Onboarding")
		self.onboarding = frappe.get_doc(
			{"doctype": "Embark Onboarding", "client_name": "Sunrise Traders", "consultant": "Administrator"}
		).insert()

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.local.conf.pop("embark_engine", None)
		frappe.db.rollback()

	def answer_interview(self, **answers):
		"""Answer every question that gets asked, following the branches."""
		while True:
			questions = api.get_interview(self.onboarding.name)["questions"]
			todo = {
				q["key"]: answers.get(q["key"], self.default_answer(q)) for q in questions if not q["answer"]
			}
			if not todo:
				return
			api.save_answers(self.onboarding.name, todo)

	@staticmethod
	def default_answer(question):
		if question["key"] == "modules":
			# The default customer wants the lot, so every step stays in play.
			return ",".join(c["value"] for c in question["choices"])
		if question["type"] == "Number":
			return "1"
		if question["type"] == "Yes / No":
			return "yes"
		return question["choices"][0]["value"]

	def upload(self, area, rows, name=None):
		"""Upload a sheet the way the portal does: a private File on the onboarding."""
		f = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": name or f"{area}.xlsx",
				"content": xlsx(rows),
				"is_private": 1,
				"attached_to_doctype": "Embark Onboarding",
				"attached_to_name": self.onboarding.name,
			}
		).insert()
		return api.attach_file(self.onboarding.name, area, f.file_url)

	def test_customer_journey(self):
		overview = api.get_overview()
		self.assertEqual(overview["name"], self.onboarding.name)
		# The company is only asked for when ERPNext has not been set up yet.
		self.assertEqual(overview["company_complete"], api.setup_done())
		self.assertEqual(overview["readiness"], 50 if api.setup_done() else 0)
		# Nothing is asked for until the interview has been answered.
		self.assertEqual(overview["steps"], [])

		# A trader, so the plan holds no workstations or operations.
		self.answer_interview(manufactures="no", subcontracts="no")
		overview = api.get_overview(self.onboarding.name)
		self.assertTrue(overview["interview"]["done"])
		self.assertEqual([s["area"] for s in overview["steps"]][:2], ["Users", "Item Groups"])

		overview = api.save_company_details(
			self.onboarding.name,
			{
				"company_name": "Sunrise Traders",
				"country": "India",
				"default_currency": "INR",
				"fiscal_year_start": "2026-04-01",
			},
		)
		self.assertTrue(overview["company_complete"])
		self.assertEqual(overview["status"], "In Progress")

		# Items first: their group is not uploaded yet, so it cannot be found.
		area = self.upload(
			"Items",
			[
				["Product List"],
				["SKU", "Product Name", "Category", "UOM", "Maintain Stock", "MRP"],
				["FG-1", "Rack", "Finished Goods", "nos", "Yes", "2,500"],
				["FG-1", "Rack copy", "Products", "Nos", "Yes", "10"],
				["RM-1", "Steel", "Products", "Kilo", "maybe", "abc"],
			],
		)
		up = area["upload"]
		mapped = {h["header"]: h["fieldname"] for h in up["headers"]}
		self.assertEqual(mapped["SKU"], "item_code")
		self.assertEqual(mapped["Category"], "item_group")
		self.assertEqual(mapped["UOM"], "stock_uom")
		codes = {(i["code"], i["fieldname"]) for i in up["issues"]}
		self.assertIn(("NOT_FOUND", "item_group"), codes)  # Finished Goods
		self.assertIn(("DUPLICATE", "item_code"), codes)
		self.assertIn(("NOT_FOUND", "stock_uom"), codes)  # Kilo
		kilo = next(i for i in up["issues"] if i["code"] == "NOT_FOUND" and i["fieldname"] == "stock_uom")
		self.assertEqual(kilo["suggestions"], ["Kg"])
		group = next(i for i in up["issues"] if i["code"] == "NOT_FOUND" and i["fieldname"] == "item_group")
		self.assertIn("Item Groups step", group["hint"])
		self.assertIn(("BAD_YESNO", "is_stock_item"), codes)
		self.assertIn(("BAD_NUMBER", "standard_rate"), codes)
		# 'nos' is the standard unit Nos, written in lower case: fixed silently.
		self.assertEqual(up["preview"]["rows"][0]["stock_uom"], "Nos")
		self.assertEqual(up["preview"]["rows"][0]["standard_rate"], 2500)

		# Uploading Item Groups re-checks Items: Finished Goods is now known.
		self.upload("Item Groups", [["Item Group", "Belongs To"], ["Finished Goods", ""]])
		up = api.get_area(self.onboarding.name, "Items")["upload"]
		self.assertNotIn(("NOT_FOUND", "item_group"), {(i["code"], i["fieldname"]) for i in up["issues"]})

		# Fix on screen: one value everywhere, and single rows.
		api.fix_value(self.onboarding.name, "Items", "stock_uom", "Kg", old_value="Kilo")
		api.fix_value(self.onboarding.name, "Items", "item_code", "FG-2", row=4)
		api.fix_value(self.onboarding.name, "Items", "is_stock_item", "Yes", row=5)
		up = api.fix_value(self.onboarding.name, "Items", "standard_rate", "90", row=5)["upload"]
		self.assertEqual(up["errors"], 0, up["issues"])

		# The file itself is untouched; fixes live beside it.
		stored = frappe.get_doc("Embark Upload", {"onboarding": self.onboarding.name, "data_area": "Items"})
		self.assertEqual(json.loads(stored.overrides)["*"]["stock_uom"], {"Kilo": "Kg"})

		self.upload("Users", [["Name", "Email"], ["Asha", "asha@sunrise.com"]])
		self.upload("Warehouses", [["Godown"], ["Main Store"], ["Shop"]])
		self.upload("Customers", [["Party Name", "GSTIN", "Mobile"], ["Acme", "27AAACS1234A1Z5", 9876543210]])

		# A column the customer's file lacks can be chosen by hand.
		area = self.upload("Suppliers", [["Vendor", "Branch Office"], ["Apex Steel", "Pune"]])
		self.assertEqual(
			{h["header"]: h["fieldname"] for h in area["upload"]["headers"]}["Branch Office"], None
		)
		area = api.set_column(self.onboarding.name, "Suppliers", "Branch Office", "address_city")
		self.assertEqual(area["upload"]["preview"]["rows"][0]["address_city"], "Pune")

		overview = api.get_overview(self.onboarding.name)
		self.assertEqual(overview["readiness"], 100, overview["steps"])
		self.assertTrue(overview["can_submit"])

		overview = api.submit_for_review(self.onboarding.name)
		self.assertEqual(overview["status"], "Submitted")
		self.assertTrue(overview["locked"])
		self.assertRaises(frappe.ValidationError, self.upload, "Users", [["Email"], ["x@y.com"]])
		api.review(self.onboarding.name, "Returned", "One more thing")
		self.upload("Users", [["Name", "Email"], ["Asha", "asha@sunrise.com"]])  # open again

		# The consultant gets ERPNext-shaped files.
		frappe.set_user("Administrator")
		api.download_prepared_data(self.onboarding.name)
		zf = zipfile.ZipFile(io.BytesIO(frappe.response.filecontent))
		items = next(n for n in zf.namelist() if n.endswith("Items.xlsx"))
		sheet = load_workbook(io.BytesIO(zf.read(items))).active
		header = [c.value for c in sheet[1]]
		self.assertIn("Item Code", header)
		self.assertEqual(sheet.max_row, 4)
		warehouses = next(n for n in zf.namelist() if n.endswith("Warehouses.xlsx"))
		self.assertEqual(load_workbook(io.BytesIO(zf.read(warehouses))).active["A2"].value, "Sunrise Traders")

		self.assertEqual(
			api.review(self.onboarding.name, "Returned", "Add your price list")["status"], "Returned"
		)
		self.assertEqual(api.get_overview()["review_notes"], "Add your price list")

	def test_remove_file(self):
		self.answer_interview()
		self.upload("Item Groups", [["Item Group"], ["Spare Parts"]])
		self.assertEqual(self.step("Item Groups")["status"], "Ready")
		area = api.remove_file(self.onboarding.name, "Item Groups")
		self.assertIsNone(area["upload"])
		self.assertEqual(self.step("Item Groups")["status"], "Not Started")
		self.assertFalse(frappe.db.exists("Embark Upload", {"onboarding": self.onboarding.name}))

		# Once sent for review, the customer can't remove files either.
		self.upload("Item Groups", [["Item Group"], ["Spare Parts"]])
		frappe.db.set_value("Embark Onboarding", self.onboarding.name, "status", "Submitted")
		self.assertRaises(frappe.ValidationError, api.remove_file, self.onboarding.name, "Item Groups")

	def step(self, area):
		return next(s for s in api.get_overview(self.onboarding.name)["steps"] if s["area"] == area)

	def test_file_must_belong_to_the_onboarding(self):
		other = frappe.get_doc(
			{
				"doctype": "File",
				"file_name": "elsewhere.xlsx",
				"content": xlsx([["Email"], ["a@b.com"]]),
				"is_private": 1,
				"attached_to_doctype": "User",
				"attached_to_name": "Administrator",
			}
		).insert()
		self.assertRaises(
			frappe.ValidationError, api.attach_file, self.onboarding.name, "Users", other.file_url
		)

	def test_start_onboarding(self):
		"""One site, one onboarding, started by a System Manager."""
		frappe.db.delete("Embark Upload")
		frappe.db.delete("Embark Onboarding")
		self.assertTrue(api.get_overview()["needs_start"])

		overview = api.start_onboarding("Yasir Traders")["overview"]
		self.assertEqual(overview["client_name"], "Yasir Traders")
		self.assertEqual(api.get_overview()["name"], overview["name"])
		# A new onboarding starts with the interview and nothing else: the steps
		# are built from the answers.
		self.assertEqual(overview["steps"], [])
		self.assertRaises(frappe.ValidationError, api.start_onboarding, "Someone Else")

	def test_standard_values_before_setup(self):
		from unittest.mock import patch

		from embark.embark.doctype.embark_upload import embark_upload
		from embark.presets import standard_values

		self.assertIn("Nos", standard_values("UOM"))
		self.assertIn("Commercial", standard_values("Customer Group", "India"))
		self.assertIn("Products", standard_values("Item Group", "India"))
		self.assertIn("India", standard_values("Territory", "India"))

		# On a site whose setup wizard has not run, those tables are empty.
		real_get_all = frappe.get_all

		def empty_masters(doctype, *args, **kwargs):
			if doctype in ("UOM", "Item Group", "Customer Group", "Supplier Group", "Territory"):
				return []
			return real_get_all(doctype, *args, **kwargs)

		columns = build_columns(frappe.get_doc("Embark Data Area", "Items"))
		with patch.object(embark_upload.frappe, "get_all", side_effect=empty_masters):
			known = embark_upload.known_values(self.onboarding.name, columns)
		self.assertEqual(known["UOM"].get("nos"), "Nos")
		self.assertEqual(known["Item Group"].get("products"), "Products")

	def test_tally_answer_offers_the_migrator(self):
		self.assertIsNone(api.get_overview(self.onboarding.name)["tally"])

		overview = api.save_answers(self.onboarding.name, {"current_system": "tally"})
		self.assertEqual(overview["tally"]["route"], "/app/tally-migrator")
		self.assertEqual(overview["tally"]["installed"], "tally_migrator" in frappe.get_installed_apps())

		# Anyone else is left alone.
		self.assertIsNone(api.save_answers(self.onboarding.name, {"current_system": "paper"})["tally"])

	def test_no_checklist_until_the_interview_is_answered(self):

		def steps():
			return [s["area"] for s in api.get_overview(self.onboarding.name)["steps"]]

		# The checklist is built from the answers, not trimmed from a standard
		# list, so a customer who has answered nothing is asked for nothing.
		self.assertEqual(steps(), [])
		self.assertRaises(frappe.DoesNotExistError, api.get_area, self.onboarding.name, "Items")

		# Half an interview is still no checklist.
		api.save_answers(self.onboarding.name, {"keeps_stock": "yes"})
		self.assertEqual(steps(), [])

		self.answer_interview()
		self.assertIn("Items", steps())

	def test_assistant_is_off_until_a_provider_is_configured(self):
		# Whatever this site is pointed at, an unconfigured one offers no chat.
		frappe.db.set_single_value("Embark Settings", {"assistant_enabled": 0, "model": "", "base_url": ""})
		frappe.clear_cache(doctype="Embark Settings")
		self.assertFalse(assistant.is_on())
		self.assertFalse(api.get_overview(self.onboarding.name)["assistant"])
		self.assertRaises(frappe.ValidationError, api.ask, "hello", self.onboarding.name)

	def test_a_tool_call_written_as_text_is_still_a_tool_call(self):
		"""A small model often types the call instead of making it."""
		from embark import llm

		written = '{"name": "save_answers", "parameters": {"answers": {"keeps_stock": "no"}}}'
		self.assertEqual(
			llm.recovered_calls(written),
			[
				{
					"id": "recovered-save_answers",
					"name": "save_answers",
					"arguments": {"answers": {"keeps_stock": "no"}},
				}
			],
		)
		# Fenced, and with a sentence around it, is the same call.
		self.assertTrue(llm.recovered_calls(f"Sure:\n```json\n{written}\n```"))
		# Plain talk is left alone.
		self.assertEqual(llm.recovered_calls("What do you sell?"), [])
		self.assertFalse(llm.looks_like_plumbing("What do you sell?"))
		self.assertTrue(llm.looks_like_plumbing(written))

	def test_the_assistant_acts_on_a_written_tool_call(self):
		from unittest.mock import patch

		frappe.local.conf["embark_engine"] = "embark"
		frappe.db.set_single_value(
			"Embark Settings",
			{
				"assistant_enabled": 1,
				"provider": "OpenAI compatible",
				"model": "stub",
				"base_url": "http://stub.invalid/v1",
			},
		)
		frappe.clear_cache(doctype="Embark Settings")

		turns = [
			{
				"text": '{"name": "save_answers", "parameters": {"answers": {"keeps_stock": "no"}}}',
				"raw": {"role": "assistant"},
				"tool_calls": [],
			},
			{"text": "Noted, no stock.", "raw": {"role": "assistant"}, "tool_calls": []},
		]
		with patch.object(assistant, "_complete", side_effect=turns):
			result = api.ask("we hold no stock", self.onboarding.name)

		self.assertEqual(result["used"], ["save_answers"])
		self.assertEqual(result["reply"], "Noted, no stock.")
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		self.assertEqual(doc.answer_map()["keeps_stock"], "no")

	def test_a_mangled_tool_call_is_still_read(self):
		"""What a small model actually writes, quotes and all."""
		from embark import llm

		mangled = '{"name":"save_answers","parameters":{"answers":"{"keeps_stock":"yes"}}}'
		calls = llm.recovered_calls(mangled)
		self.assertEqual(calls[0]["name"], "save_answers")
		self.assertEqual(calls[0]["arguments"], {"answers": {"keeps_stock": "yes"}})

	def test_only_what_was_said_goes_back_to_the_model(self):
		"""The panel draws tool steps beside the talk; they are not conversation."""
		history = [
			{"role": "user", "content": "hello"},
			{"role": "step", "step": {"tool": "open_questions"}},
			{"role": "assistant", "content": "What do you sell?"},
			{"role": "change", "set": {"name": "CHG-0001"}},
			{"role": "assistant", "content": ""},
		]
		self.assertEqual(
			assistant._conversation(history),
			[{"role": "user", "content": "hello"}, {"role": "assistant", "content": "What do you sell?"}],
		)
		self.assertEqual(assistant._conversation(None), [])

	def test_assistant_answers_the_interview_and_nothing_else(self):
		from unittest.mock import patch

		frappe.local.conf["embark_engine"] = "embark"
		frappe.db.set_single_value(
			"Embark Settings",
			{
				"assistant_enabled": 1,
				"provider": "OpenAI compatible",
				"model": "stub",
				"base_url": "http://stub.invalid/v1",
			},
		)
		frappe.clear_cache(doctype="Embark Settings")
		self.assertTrue(assistant.is_on())

		# The model looks at what is open, saves what the customer described,
		# then answers in words.
		turns = [
			{
				"text": "",
				"raw": {"role": "assistant"},
				"tool_calls": [{"id": "1", "name": "open_questions", "arguments": {}}],
			},
			{
				"text": "",
				"raw": {"role": "assistant"},
				"tool_calls": [
					{
						"id": "2",
						"name": "save_answers",
						"arguments": {"answers": {"keeps_stock": "no", "business_type": "services"}},
					}
				],
			},
			{"text": "Noted: services, no stock.", "raw": {"role": "assistant"}, "tool_calls": []},
		]
		with patch.object(assistant, "_complete", side_effect=turns):
			result = api.ask("We're a services firm, we hold no stock.", self.onboarding.name)

		self.assertEqual(result["reply"], "Noted: services, no stock.")
		self.assertEqual(result["used"], ["open_questions", "save_answers"])
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		self.assertEqual(doc.answer_map()["keeps_stock"], "no")
		# It went through save_answers, so the same validation applies to it.
		self.assertEqual(
			assistant._run_tool(self.onboarding.name, "save_answers", {"answers": {"nope": "yes"}}).get(
				"error"
			)
			is not None,
			True,
		)

	def test_the_plan_is_generated_from_the_answers(self):

		def plan():
			return {line["key"]: line for line in api.get_overview(self.onboarding.name)["plan"]}

		# No interview, no plan — the settings and training are not a standard list.
		self.assertEqual(plan(), {})

		self.answer_interview(keeps_stock="no", manufactures="no", tax_registered="no")
		lines = plan()
		self.assertNotIn("batches", lines)
		self.assertNotIn("costing_method", lines)
		self.assertNotIn("training_stock", lines)
		self.assertNotIn("tax_setup", lines)
		self.assertIn("training_selling", lines)

		# Keeping stock brings the stock side of the plan with it, and says why.
		api.save_answers(
			self.onboarding.name,
			{"keeps_stock": "yes", "tracks_batches": "yes", "tax_registered": "yes"},
		)
		self.answer_interview()
		lines = plan()
		self.assertEqual(lines["batches"]["kind"], "Setting")
		self.assertEqual(lines["batches"]["because"], "you track batches or expiry dates")
		self.assertEqual(lines["costing_method"]["kind"], "Decision")
		self.assertIn("training_stock", lines)
		self.assertIn("tax_setup", lines)

		# A line the consultant has settled survives a change of mind.
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		for row in doc.tasks:
			if row.setup_task == "batches":
				row.status = "Done"
		doc.save()
		api.save_answers(self.onboarding.name, {"tracks_batches": "no"})
		self.assertEqual(plan()["batches"]["status"], "Done")

	def test_interview_shapes_the_onboarding(self):

		def steps():
			return [s["area"] for s in api.get_overview(self.onboarding.name)["steps"]]

		def item_columns():
			return {c["fieldname"] for c in api.get_area(self.onboarding.name, "Items")["columns"]}

		# A business with no stock is not asked for warehouses or stock columns.
		self.answer_interview(keeps_stock="no", business_type="services")
		overview = api.get_overview(self.onboarding.name)
		self.assertNotIn("Warehouses", steps())
		self.assertEqual(item_columns() & {"has_batch_no", "has_serial_no", "is_stock_item"}, set())
		self.assertTrue(overview["interview"]["done"])
		# The stock questions are no longer asked, so they are not counted either.
		asked = [q["key"] for q in api.get_interview(self.onboarding.name)["questions"]]
		self.assertNotIn("tracks_batches", asked)

		# The template follows the same columns.
		api.download_template("Items", self.onboarding.name)
		headers = [c.value for c in load_workbook(io.BytesIO(frappe.response.filecontent)).active[1]]
		self.assertNotIn("Batch Tracked", headers)

		# Saying yes again brings the step and its questions back.
		api.save_answers(self.onboarding.name, {"keeps_stock": "yes"})
		self.assertIn("Warehouses", steps())
		self.assertIn(
			"tracks_batches", [q["key"] for q in api.get_interview(self.onboarding.name)["questions"]]
		)

		# A step that already holds an upload is never taken away.
		self.upload("Warehouses", [["Godown"], ["Main Store"]])
		api.save_answers(self.onboarding.name, {"keeps_stock": "no"})
		self.assertIn("Warehouses", steps())

		self.assertRaises(
			frappe.ValidationError, api.save_answers, self.onboarding.name, {"keeps_stock": "maybe"}
		)
		self.assertRaises(frappe.ValidationError, api.save_answers, self.onboarding.name, {"nope": "yes"})

	def test_a_missing_column_is_one_problem_not_one_per_row(self):
		"""A customer with 3,000 items must not be shown 3,000 errors for one column."""
		self.answer_interview()
		rows = [["Item Code", "Item Name", "Item Group"]]
		rows += [[f"SKU-{i:04d}", f"Part {i}", "Products"] for i in range(50)]
		area = self.upload("Items", rows)
		up = area["upload"]

		codes = [i["code"] for i in up["issues"]]
		self.assertIn("COLUMN_MISSING", codes)
		# One line about the unit column, not one per row.
		self.assertLessEqual(codes.count("REQUIRED_MISSING"), 1)
		self.assertLess(up["errors"], 10, up["issues"][:3])

	def test_rows_can_be_typed_instead_of_uploaded(self):
		"""Three warehouses do not deserve a spreadsheet."""
		self.answer_interview()
		area = api.enter_rows(
			self.onboarding.name,
			"Warehouses",
			[
				{"warehouse_name": "Main Store"},
				{"warehouse_name": "Godown 2"},
				{"warehouse_name": ""},  # a blank line is not a row
			],
		)
		up = area["upload"]
		self.assertEqual(up["rows"], 2)
		self.assertEqual(up["errors"], 0, up["issues"])
		self.assertEqual(self.step("Warehouses")["status"], "Ready")

		# It became a real sheet, so everything downstream works on it.
		self.assertTrue(up["file_name"].endswith(".xlsx"))
		self.assertEqual({r["warehouse_name"] for r in up["preview"]["rows"]}, {"Main Store", "Godown 2"})

		self.assertRaises(frappe.ValidationError, api.enter_rows, self.onboarding.name, "Warehouses", [])
		self.assertRaises(
			frappe.ValidationError, api.enter_rows, self.onboarding.name, "Warehouses", [{}] * 300
		)

	def test_standard_masters_count_as_known(self):
		"""Standard Selling is ERPNext's own price list, not a typo."""
		self.answer_interview(price_lists="yes")
		area = self.upload(
			"Item Prices",
			[
				["Item Code", "Price List", "Rate"],
				["SKU-1", "Standard Selling", "100"],
				["SKU-1", "Standard Buying", "80"],
				["SKU-1", "Wholesael", "90"],  # a real typo, and no list of that name
			],
		)
		up = area["upload"]
		unknown = [i for i in up["issues"] if i["code"] == "NOT_FOUND" and i["fieldname"] == "price_list"]
		self.assertEqual([i["value"] for i in unknown], ["Wholesael"], up["issues"])
		# And a suggestion is never the value already typed.
		for issue in up["issues"]:
			self.assertNotIn(issue["value"], issue.get("suggestions", []))

	def test_template_download(self):
		self.answer_interview()
		api.download_template("Customers", self.onboarding.name)
		wb = load_workbook(io.BytesIO(frappe.response.filecontent))
		self.assertEqual(wb.sheetnames[1], "How to fill")
		self.assertEqual(wb.active["A1"].value, "Customer Name *")

	def test_approving_creates_the_records_in_erpnext(self):
		"""The point of the whole thing: approved data becomes real records."""
		if not api.setup_done():
			self.skipTest("no company on this site yet")
		self.answer_interview()
		api.enter_rows(
			self.onboarding.name,
			"Warehouses",
			[{"warehouse_name": "Embark Test Store"}, {"warehouse_name": "Embark Test Godown"}],
		)
		self.upload(
			"Customers",
			[["Customer Name", "Customer Group"], ["Embark Test Customer", "Commercial"]],
		)
		result = api.review(self.onboarding.name, "Approved")
		summary = result["summary"]
		self.assertIsNotNone(summary, "approving did not create anything")

		warehouses = next(s for s in summary["steps"] if s["area"] == "Warehouses")
		self.assertEqual(warehouses["created"], 2, warehouses)
		self.assertTrue(frappe.db.exists("Warehouse", {"warehouse_name": "Embark Test Store"}))
		self.assertTrue(frappe.db.exists("Customer", {"customer_name": "Embark Test Customer"}))

		# The company came from the onboarding, not from a guess.
		self.assertEqual(
			frappe.db.get_value("Warehouse", {"warehouse_name": "Embark Test Store"}, "company"),
			summary["company"],
		)

		# A step nobody filled in is reported, not failed.
		empty = [s for s in summary["steps"] if s["reason"]]
		for step in empty:
			self.assertEqual(step["created"], 0)

		# And it is safe to press twice: the second run creates nothing.
		again = api.create_in_erpnext(self.onboarding.name)["summary"]
		self.assertEqual(again["created"], 0, again)
		self.assertEqual(again["skipped"], 3, again)
		self.assertEqual(again["failed"], 0, again)

	def test_nothing_is_created_from_a_step_with_errors(self):
		"""Bad data is never half-imported."""
		if not api.setup_done():
			self.skipTest("no company on this site yet")
		self.answer_interview()
		self.upload(
			"Warehouses",
			# The same store twice: an error, and neither row is guessed at.
			[["Warehouse Name"], ["Embark Bad Store"], ["Embark Bad Store"]],
		)
		frappe.db.set_value("Embark Onboarding", self.onboarding.name, "status", "Submitted")
		summary = api.create_in_erpnext(self.onboarding.name)["summary"]
		warehouses = next(s for s in summary["steps"] if s["area"] == "Warehouses")
		self.assertIn("fix", warehouses["reason"])
		self.assertEqual(warehouses["created"], 0)
		self.assertFalse(frappe.db.exists("Warehouse", {"warehouse_name": "Embark Bad Store"}))

	def test_the_modules_chosen_decide_the_whole_checklist(self):
		"""The first question is which parts of ERPNext are wanted."""
		first = api.get_interview(self.onboarding.name)["questions"][0]
		self.assertEqual(first["key"], "modules")

		# Accounting and HR only: no stock, no selling, no making anything.
		self.answer_interview(modules="accounts,hr")
		asked = {q["key"] for q in api.get_interview(self.onboarding.name)["questions"]}
		self.assertNotIn("keeps_stock", asked)
		self.assertNotIn("manufactures", asked)
		self.assertNotIn("sells_on_orders", asked)

		steps = {s["area"] for s in api.get_overview(self.onboarding.name)["steps"]}
		self.assertIn("Users", steps)
		self.assertNotIn("Items", steps)
		self.assertNotIn("Warehouses", steps)
		self.assertNotIn("Workstations", steps)
		self.assertNotIn("Sales People", steps)

		plan = {t["key"] for t in api.get_overview(self.onboarding.name)["plan"]}
		self.assertIn("hr_setup", plan)
		self.assertIn("accounts_setup", plan)
		self.assertNotIn("manufacturing", plan)

	def test_choosing_manufacturing_asks_for_the_shop_floor(self):
		self.answer_interview(modules="stock,manufacturing", manufactures="yes")
		steps = {s["area"] for s in api.get_overview(self.onboarding.name)["steps"]}
		self.assertIn("Workstations", steps)
		self.assertIn("Items", steps)
		self.assertIn("Warehouses", steps)
		# Nobody asked for selling, so none of its questions were put.
		asked = {q["key"] for q in api.get_interview(self.onboarding.name)["questions"]}
		self.assertNotIn("price_lists", asked)

	# ── The chat as a way of doing the work ───────────────────────────────

	def chat_on(self, engine="embark"):
		"""Turn the chat on, pinned to the engine under test.

		Flow is what runs the chat on a real site; these tests exercise Embark's
		own loop through a stubbed provider, so they say so rather than letting
		the engine reach for a model that is not there.
		"""
		frappe.local.conf["embark_engine"] = engine
		frappe.db.set_single_value(
			"Embark Settings",
			{
				"assistant_enabled": 1,
				"provider": "OpenAI compatible",
				"model": "stub",
				"base_url": "http://stub.invalid/v1",
			},
		)
		frappe.clear_cache(doctype="Embark Settings")

	@staticmethod
	def says(text="", calls=None):
		return {"text": text, "raw": {"role": "assistant"}, "tool_calls": calls or []}

	@staticmethod
	def call(name, **arguments):
		return {"id": f"call_{name}", "name": name, "arguments": arguments}

	def test_the_chat_only_offers_the_tools_that_fit_where_you_are(self):
		"""Twelve choices is how a model goes wrong. It gets the ones that apply."""
		names = {t["name"] for t in assistant.tools_for(self.onboarding.name)}
		self.assertIn("open_questions", names)
		self.assertNotIn("check_step", names, "data tools before the interview is answered")
		self.assertNotIn("send_for_review", names)

		self.answer_interview()
		names = {t["name"] for t in assistant.tools_for(self.onboarding.name)}
		self.assertIn("check_step", names)
		self.assertIn("fix", names)
		self.assertNotIn("open_questions", names, "nothing left to ask")
		# Sending is offered only when there is something to send.
		self.assertNotIn("send_for_review", names)

	def test_the_chat_checks_a_sheet_and_fixes_it(self):
		"""The whole point: the work happens in the conversation."""
		from unittest.mock import patch

		self.chat_on()
		self.answer_interview(price_lists="yes")
		self.upload(
			"Item Prices",
			[
				["Item Code", "Price List", "Rate"],
				["SKU-1", "Standard Selling", "100"],
				["SKU-2", "Wholesael", "90"],
			],
		)

		turns = [
			self.says(calls=[self.call("check_step", area="Item Prices")]),
			self.says("One price list is spelt wrong. Shall I make it Standard Selling?"),
		]
		with patch.object(assistant, "_complete", side_effect=turns):
			result = api.ask("is my price list alright?", self.onboarding.name)

		self.assertEqual(result["used"], ["check_step"])
		# What it did is shown in words, never as a tool name.
		self.assertEqual(result["steps"][0]["label"], "Checked your Item Prices")
		self.assertFalse(result["steps"][0]["failed"])
		before = self.step("Item Prices")["errors"]

		turns = [
			self.says(
				calls=[
					self.call(
						"fix",
						area="Item Prices",
						column="price_list",
						old_value="Wholesael",
						value="Standard Selling",
					)
				]
			),
			self.says("Fixed."),
		]
		with patch.object(assistant, "_complete", side_effect=turns):
			result = api.ask("yes please", self.onboarding.name)

		# The misspelt price list is gone; the unknown item codes are not its
		# business, and are still there to be dealt with on their own step.
		self.assertEqual(self.step("Item Prices")["errors"], before - 1, result["steps"])
		issues = api.get_area(self.onboarding.name, "Item Prices")["upload"]["issues"]
		self.assertNotIn("Wholesael", [i["value"] for i in issues])
		self.assertIn("Standard Selling", result["steps"][0]["label"])

	def test_the_chat_cannot_reach_a_step_that_is_not_yours(self):
		self.chat_on()
		self.answer_interview()
		result = assistant._run_tool(self.onboarding.name, "check_step", {"area": "Not A Step"})
		self.assertIn("error", result)

	def test_what_the_customer_said_is_written_down_even_if_the_model_forgets(self):
		"""A model that asks well and saves nothing leaves them repeating themselves."""
		from unittest.mock import patch

		self.chat_on()
		turns = [
			# It answers in words and never calls save_answers…
			self.says("Right, no stock. What do you use today?"),
			# …so the catch-up reads the same sentence, on its own, as JSON.
			self.says('{"keeps_stock": "no"}'),
		]
		with patch.object(assistant, "_complete", side_effect=turns):
			result = api.ask("we hold no stock at all", self.onboarding.name)

		self.assertIn("save_answers", result["used"])
		self.assertEqual(self.onboarding.reload().answer_map()["keeps_stock"], "no")

	def test_the_catch_up_saves_only_real_answers(self):
		"""Whatever the model returns, only a real choice of a real question is kept."""
		open_now = [
			{
				"key": "keeps_stock",
				"question": "?",
				"type": "Yes / No",
				"choices": ["yes", "no"],
				"allow_not_sure": True,
			},
			{
				"key": "business_type",
				"question": "?",
				"type": "Several choices",
				"choices": ["trading", "retail"],
				"allow_not_sure": True,
			},
		]
		self.assertEqual(
			assistant._real_answers(
				{
					"keeps_stock": "yes",
					"business_type": "trading,retail,banking",
					"made_up_question": "yes",
					"keeps_stock_extra": "yes",
				},
				open_now,
			),
			{"keeps_stock": "yes", "business_type": "trading,retail"},
		)

	def test_a_file_sent_in_the_chat_lands_on_the_right_step(self):
		self.answer_interview()
		steps = [s["area"] for s in api.get_overview(self.onboarding.name)["steps"]]
		self.assertEqual(assistant._best_step("customer list final.xlsx", steps), "Customers")
		self.assertEqual(assistant._best_step("our-warehouses.xlsx", steps), "Warehouses")
		# A name that says nothing is not guessed at: the chat asks instead.
		self.assertIsNone(assistant._best_step("book1.xlsx", steps))

	# ── Configuring ERPNext from the plan ─────────────────────────────────

	def test_the_plan_configures_erpnext_itself(self):
		"""The point of the tool: the setting is written, not written down."""
		self.answer_interview()
		before = frappe.db.get_single_value("Selling Settings", "so_required")
		frappe.db.set_single_value("Selling Settings", "so_required", "No")
		frappe.clear_cache(doctype="Selling Settings")

		shown = api.configure_preview(self.onboarding.name, "sales_flow")["changes"]
		self.assertEqual(shown[0]["doctype"], "Selling Settings")
		self.assertEqual(shown[0]["lines"][0]["before"], "No")
		self.assertEqual(shown[0]["lines"][0]["after"], "Yes")

		out = api.configure_step(self.onboarding.name, "sales_flow")["result"]
		self.assertFalse(out["already"])
		self.assertEqual(frappe.db.get_single_value("Selling Settings", "so_required"), "Yes")

		# The line is marked done, and says what did it.
		line = next(t for t in api.get_overview(self.onboarding.name)["plan"] if t["key"] == "sales_flow")
		self.assertEqual(line["status"], "Done")
		self.assertIn("Embark", line["notes"])

		# It is undoable, because it went through the same change set Studio uses.
		self.assertTrue(frappe.db.exists("Embark Change Set", out["change_set"]))
		frappe.get_doc("Embark Change Set", out["change_set"]).undo()
		self.assertEqual(frappe.db.get_single_value("Selling Settings", "so_required"), "No")
		frappe.db.set_single_value("Selling Settings", "so_required", before)

	def test_applying_a_second_time_changes_nothing(self):
		self.answer_interview()
		frappe.db.set_single_value("Selling Settings", "so_required", "Yes")
		frappe.clear_cache(doctype="Selling Settings")
		out = api.configure_step(self.onboarding.name, "sales_flow")["result"]
		self.assertTrue(out["already"], out)
		self.assertIsNone(out["change_set"])

	def test_configure_all_does_every_line_it_can(self):
		from embark import configure

		self.answer_interview()
		out = api.configure_all(self.onboarding.name)
		self.assertEqual(out["failed"], [])
		self.assertTrue(out["configured"] or out["already"])

		# Every line Embark claims it can do is now done; the rest stay for a person.
		plan = api.get_overview(self.onboarding.name)["plan"]
		for line in plan:
			if configure.can_apply(line["key"]):
				self.assertEqual(line["status"], "Done", line["key"])

		# And the switches really moved.
		self.assertEqual(frappe.db.get_single_value("Buying Settings", "po_required"), "Yes")
		self.assertEqual(
			frappe.db.get_single_value("Stock Settings", "enable_serial_and_batch_no_for_item"), 1
		)

	def test_a_line_with_no_recipe_is_left_for_a_person(self):
		from embark import configure

		self.assertFalse(configure.can_apply("document_numbering"))
		self.assertFalse(configure.can_apply("training_selling"))
		self.answer_interview()
		line = next(
			t for t in api.get_overview(self.onboarding.name)["plan"] if t["key"] == "document_numbering"
		)
		self.assertFalse(line["can_apply"])
		self.assertEqual(line["status"], "Planned")

	# ── Ask Embark on Frappe Flow ─────────────────────────────────────────

	def test_ask_embark_runs_on_flow_when_it_can(self):
		"""The chat is the same chat; Flow is what runs it."""
		from unittest.mock import patch

		from embark import assistant_flow

		self.chat_on()
		self.answer_interview()

		class Call:
			name = "where_am_i"
			arguments: typing.ClassVar[dict] = {}
			result = '{"customer": "Sunrise Traders"}'

		class Result:
			output = "You are two steps from done."
			tool_calls: typing.ClassVar[list] = [Call()]

		class Engine:
			def __init__(self, **kw):
				Engine.tools = kw["tools"]
				Engine.instructions = kw["instructions"]

			def run(self, messages):
				Engine.messages = messages
				return Result()

		with (
			patch.object(assistant_flow, "available", return_value=True),
			patch.object(assistant_flow, "model", return_value=object()),
			patch("embark.vendor.flow.Agent", Engine),
		):
			out = api.ask("where am I?", self.onboarding.name)

		self.assertEqual(out["engine"], "flow")
		self.assertEqual(out["reply"], "You are two steps from done.")
		# What it did is still shown in words, not as a tool name.
		self.assertEqual(out["steps"][0]["label"], "Checked where you are")

		# Flow was handed typed tools, and only the ones this stage offers.
		names = {t.name for t in Engine.tools}
		self.assertIn("check_step", names)
		self.assertNotIn("open_questions", names, "the interview is answered")
		# Their schemas come from the function signatures, not a hand-written dict.
		step = next(t for t in Engine.tools if t.name == "step_help")
		self.assertIn("area", step.to_dict()["function"]["parameters"]["properties"])

	def test_the_chat_falls_back_when_flow_cannot_run(self):
		"""A site without litellm keeps its chat, and a site can pin either engine."""
		from unittest.mock import patch

		from embark import assistant_flow

		self.chat_on(engine="flow")

		# Flow talks to providers through litellm. Without it, Embark's own loop.
		with patch.dict("sys.modules", {"litellm": None}):
			self.assertFalse(assistant_flow.available())

		# And whatever is installed, a site can say which engine it wants.
		frappe.local.conf["embark_engine"] = "embark"
		self.assertFalse(assistant_flow.available())

	def test_a_provider_that_cannot_be_reached_says_why(self):
		"""'Could not be reached' is the same message for a wrong key, a wrong
		address and a model that does not exist. Whoever set it up needs the
		difference; the customer does not."""
		import requests

		self.chat_on()
		frappe.db.set_single_value(
			"Embark Settings", {"model": "ghost-model", "base_url": "http://nowhere.invalid/v1"}
		)
		frappe.clear_cache(doctype="Embark Settings")

		class Response:
			text = '{"error": "model \'ghost-model\' not found"}'

		failure = requests.RequestException("404 Client Error")
		failure.response = Response()

		said = assistant._unreachable(failure)
		self.assertIn("still work", said)
		self.assertIn("ghost-model", said, "staff are told which model")
		self.assertIn("nowhere.invalid", said, "and where it was asked")
		self.assertIn("not found", said, "and what the provider itself said")

		# The customer gets the plain sentence and nothing to puzzle over.
		with patch("embark.permissions.is_staff", return_value=False):
			plain = assistant._unreachable(failure)
		self.assertNotIn("ghost-model", plain)
		self.assertNotIn("nowhere.invalid", plain)

	def test_a_customer_arrives_with_somewhere_to_invoice_and_someone_to_write_to(self):
		"""A party with no address cannot be invoiced, and ERPNext keeps the
		email and phone on a Contact — the fields on the customer are read-only
		copies of it."""
		if not api.setup_done():
			self.skipTest("no company on this site yet")
		self.answer_interview()
		self.upload(
			"Customers",
			[
				[
					"Customer Name",
					"Tax ID",
					"Contact",
					"Email",
					"Mobile",
					"Address",
					"City",
					"State",
					"PIN / Postal Code",
				],
				[
					"Northwind Traders",
					"27AAACS1234A1Z5",
					"Aisha Khan",
					"aisha@northwind.example",
					"9876543210",
					"14 Nehru Road",
					"Pune",
					"Maharashtra",
					"411001",
				],
			],
		)
		summary = api.create_in_erpnext(self.onboarding.name)["summary"]
		step = next(s for s in summary["steps"] if s["area"] == "Customers")
		self.assertEqual(step["created"], 1, step)
		self.assertEqual(step["also"], {"Address": 1, "Contact": 1}, step)

		self.assertEqual(frappe.db.get_value("Customer", "Northwind Traders", "tax_id"), "27AAACS1234A1Z5")

		address = frappe.db.get_value(
			"Dynamic Link",
			{"link_doctype": "Customer", "link_name": "Northwind Traders", "parenttype": "Address"},
			"parent",
		)
		self.assertTrue(address, "no address to invoice them at")
		self.assertEqual(
			frappe.db.get_value("Address", address, ["address_line1", "city", "state", "pincode"]),
			("14 Nehru Road", "Pune", "Maharashtra", "411001"),
		)

		contact = frappe.db.get_value(
			"Dynamic Link",
			{"link_doctype": "Customer", "link_name": "Northwind Traders", "parenttype": "Contact"},
			"parent",
		)
		self.assertTrue(contact, "nobody to write to")
		self.assertEqual(frappe.db.get_value("Contact", contact, "first_name"), "Aisha")
		# ERPNext copies them onto the customer from the primary contact.
		self.assertEqual(
			frappe.db.get_value("Contact Email", {"parent": contact}, "email_id"),
			"aisha@northwind.example",
		)

		# Pressing again adds no second address and no second contact.
		again = api.create_in_erpnext(self.onboarding.name)["summary"]
		self.assertEqual(again["also"], {}, again)

	def test_a_party_with_no_contact_details_still_imports(self):
		"""Most sheets have a name and nothing else. That is still a customer."""
		if not api.setup_done():
			self.skipTest("no company on this site yet")
		self.answer_interview()
		self.upload("Customers", [["Customer Name"], ["Bare Name Only"]])
		summary = api.create_in_erpnext(self.onboarding.name)["summary"]
		step = next(s for s in summary["steps"] if s["area"] == "Customers")
		self.assertEqual(step["created"], 1)
		self.assertEqual(step["also"], {})
		self.assertTrue(frappe.db.exists("Customer", "Bare Name Only"))

	def test_the_tax_rates_they_charge_become_real_templates(self):
		"""ERPNext ships no tax templates. Without them nobody can raise a
		correct invoice, so a standard implementation is not done until they
		exist."""
		if not api.setup_done():
			self.skipTest("no company on this site yet")
		self.answer_interview(tax_registered="yes", tax_rates="18")
		company = frappe.get_all("Company", pluck="name", order_by="creation asc")[0]
		abbr = frappe.db.get_value("Company", company, "abbr")

		# On a site where this step has already run, the recipe plans nothing —
		# which is the idempotence checked at the end of this test. What the
		# site holds afterwards is the same either way, and that is the point.
		fresh = not frappe.db.exists("Account", f"Output Tax 18% - {abbr}")
		shown = api.configure_preview(self.onboarding.name, "tax_setup")["changes"]
		if fresh:
			self.assertTrue(shown, "nothing planned for the rate they charge")

		api.configure_step(self.onboarding.name, "tax_setup")

		for side in ("Output", "Input"):
			account = f"{side} Tax 18% - {abbr}"
			self.assertTrue(frappe.db.exists("Account", account), account)
			self.assertEqual(frappe.db.get_value("Account", account, "account_type"), "Tax")

		for doctype, head in (
			("Sales Taxes and Charges Template", f"Output Tax 18% - {abbr}"),
			("Purchase Taxes and Charges Template", f"Input Tax 18% - {abbr}"),
		):
			name = frappe.db.get_value(doctype, {"title": "Tax 18%", "company": company})
			self.assertTrue(name, doctype)
			row = frappe.get_doc(doctype, name).taxes[0]
			self.assertEqual((row.account_head, row.rate, row.charge_type), (head, 18.0, "On Net Total"))

		item_tax = frappe.db.get_value("Item Tax Template", {"title": "Tax 18%", "company": company})
		self.assertTrue(item_tax, "no item tax group to put on the items")
		self.assertEqual(
			sorted(r.tax_rate for r in frappe.get_doc("Item Tax Template", item_tax).taxes), [18.0, 18.0]
		)

		# Nothing is planned twice.
		self.assertEqual(api.configure_preview(self.onboarding.name, "tax_setup")["changes"], [])

	def test_no_tax_rates_means_nothing_is_invented(self):
		"""A business that is not registered gets no templates and no accounts."""
		self.answer_interview(tax_registered="no")
		from embark import configure

		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		self.assertEqual(configure.plan_changes(doc, "tax_setup"), [])


class TestSetupWizard(TestCase):
	"""The wizard: three stages, questions that follow from the modules, and
	steps that write real records.

	bench --site <site> run-tests --module embark.tests.test_onboarding
	"""

	def setUp(self):
		frappe.set_user("Administrator")
		frappe.db.delete("Embark Onboarding")
		self.onboarding = frappe.get_doc(
			{"doctype": "Embark Onboarding", "client_name": "Sunrise Traders", "consultant": "Administrator"}
		).insert()

	def tearDown(self):
		frappe.set_user("Administrator")
		frappe.db.rollback()

	def board(self):
		return wizard.board(frappe.get_doc("Embark Onboarding", self.onboarding.name))

	def answer(self, key, value):
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		wizard.save_answer(doc, key, value)

	def task(self, board, title):
		return next((t for t in board["tasks"] if t["title"] == title), None)

	def test_the_three_stages_are_in_order(self):
		"""Who the company is, what it runs, then what that calls for."""
		stages = [s["key"] for s in self.board()["stages"]]
		self.assertEqual(stages, ["Company", "Modules", "Setup"])

	def test_a_stage_counts_only_its_own_questions(self):
		board = self.board()
		for stage in board["stages"]:
			mine = [t for t in board["tasks"] if t["stage"] == stage["key"]]
			self.assertEqual(stage["total"], sum(t["total"] for t in mine), stage["key"])
			self.assertEqual(stage["answered"], sum(t["answered"] for t in mine), stage["key"])

	def test_the_modules_decide_what_is_asked(self):
		"""The whole point of asking for modules second."""
		self.answer("modules", "accounts")
		without = self.board()
		self.assertIsNone(self.task(without, "Set up warehouses"))

		self.answer("modules", "accounts,stock")
		with_stock = self.board()
		self.assertIsNotNone(self.task(with_stock, "Set up warehouses"))
		self.assertGreater(with_stock["total"], without["total"])

	def test_dropping_a_module_hides_its_answers_but_keeps_them(self):
		"""Going back is safe: the answers are still there if they return."""
		self.answer("modules", "accounts,stock")
		self.answer("additional_warehouse_name", "Main Store")
		self.answer("modules", "accounts")
		self.assertIsNone(self.task(self.board(), "Set up warehouses"))

		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		kept = {row.question: row.value for row in doc.wizard_answers}
		self.assertEqual(kept.get("additional_warehouse_name"), "Main Store")

	def test_a_form_of_answers_is_saved_in_one_go(self):
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		wizard.save_answers(doc, {"company_name": "Sunrise Traders Pvt Ltd", "company_abbr": "STPL"})
		kept = {row.question: row.value for row in doc.wizard_answers}
		self.assertEqual(kept["company_name"], "Sunrise Traders Pvt Ltd")
		self.assertEqual(kept["company_abbr"], "STPL")

	def test_an_answer_to_a_question_that_does_not_exist_is_refused(self):
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		self.assertRaises(frappe.ValidationError, wizard.save_answers, doc, {"not_a_question": "x"})

	def test_one_step_can_write_more_than_one_record(self):
		"""A step is a topic, and a topic can reach several records."""
		task = frappe.get_doc(
			{
				"doctype": "Embark Wizard Task",
				"task_name": "Test step: two records",
				"sequence": 999,
				"stage": "Setup",
				"description": "Test step.",
			}
		).insert(ignore_permissions=True)
		for key, label, doctype, field, kind in (
			("test_term_name", "What do you call it?", "Payment Term", "payment_term_name", "Text"),
			("test_term_days", "How many days?", "Payment Term", "credit_days", "Number"),
			("test_credit_limit", "How much credit?", "Company", "credit_limit", "Number"),
		):
			frappe.get_doc(
				{
					"doctype": "Embark Wizard Question",
					"question_key": key,
					"task": task.name,
					"label": label,
					"target_doctype": doctype,
					"fieldname": field,
					"answer_type": kind,
					"sequence": 10,
				}
			).insert(ignore_permissions=True)

		self.answer("test_term_name", "Net 30")
		self.answer("test_term_days", "30")
		self.answer("test_credit_limit", "500000")

		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		plans = wizard._records(doc, task.name)
		self.assertEqual(
			sorted(p["doctype"] for p in plans), ["Company", "Payment Term"], "both records planned"
		)
		term = next(p for p in plans if p["doctype"] == "Payment Term")
		self.assertEqual(term["values"]["payment_term_name"], "Net 30")
		self.assertEqual(term["values"]["credit_days"], "30")

		shown = wizard.check(doc, task.name)
		self.assertTrue(shown["ready"])
		self.assertEqual(len(shown["records"]), 2, shown["records"])
		self.assertIn("Payment Term", shown["about_to"])
		self.assertIn("Company", shown["about_to"])

	def test_a_question_with_no_field_decides_rather_than_writes(self):
		"""'Do you inspect goods?' is an answer, not a value to write."""
		task = frappe.get_doc(
			{
				"doctype": "Embark Wizard Task",
				"task_name": "Test step: a decision",
				"sequence": 998,
				"stage": "Setup",
				"target_doctype": "Quality Inspection Template",
				"description": "Test step.",
			}
		).insert(ignore_permissions=True)
		frappe.get_doc(
			{
				"doctype": "Embark Wizard Question",
				"question_key": "test_does_qi",
				"task": task.name,
				"label": "Do you inspect goods?",
				"answer_type": "Yes / No",
				"sequence": 10,
			}
		).insert(ignore_permissions=True)
		frappe.get_doc(
			{
				"doctype": "Embark Wizard Question",
				"question_key": "test_qi_template",
				"task": task.name,
				"label": "What is the checklist called?",
				"fieldname": "quality_inspection_template_name",
				"answer_type": "Text",
				"applies_when": "test_does_qi == yes",
				"sequence": 20,
			}
		).insert(ignore_permissions=True)

		self.answer("test_does_qi", "No")
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		self.assertEqual(wizard._records(doc, task.name), [], "nothing to write, and nothing asked")

		self.answer("test_does_qi", "Yes")
		self.answer("test_qi_template", "Kraft Paper Incoming Check")
		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		plans = wizard._records(doc, task.name)
		self.assertEqual(len(plans), 1)
		self.assertEqual(
			plans[0]["values"], {"quality_inspection_template_name": "Kraft Paper Incoming Check"}
		)

	def test_a_child_table_answer_rides_on_the_record_that_holds_it(self):
		"""How an Asset Category depreciates is a row on that category, not a
		record of its own."""
		if not frappe.db.exists("DocType", "Asset Finance Book"):
			self.skipTest("this site has no fixed assets")
		self.answer("modules", "assets")
		self.answer("tracks_fixed_assets", "Yes")
		self.answer("asset_category_name", "Plant and Machinery")
		self.answer("depreciation_method", "Written Down Value")
		self.answer("total_number_of_depreciations", "15")
		self.answer("frequency_of_depreciation", "12")

		doc = frappe.get_doc("Embark Onboarding", self.onboarding.name)
		plans = wizard._records(doc, "Set up fixed assets")
		category = next(p for p in plans if p["doctype"] == "Asset Category")
		self.assertEqual(category["values"]["asset_category_name"], "Plant and Machinery")
		# The three depreciation answers are one row on the category, and there
		# is no plan to create a child record on its own.
		self.assertEqual(
			category["values"]["finance_books"],
			[
				{
					"depreciation_method": "Written Down Value",
					"total_number_of_depreciations": "15",
					"frequency_of_depreciation": "12",
				}
			],
		)
		self.assertNotIn("Asset Finance Book", [p["doctype"] for p in plans])

	def test_a_question_needing_an_app_that_is_not_installed_is_not_asked(self):
		"""Payroll questions arrive with Frappe HR, not before it."""
		if frappe.db.exists("DocType", "Payroll Settings"):
			self.skipTest("Frappe HR is installed on this site")
		self.answer("modules", "hr")
		self.assertIsNone(self.task(self.board(), "Set payroll rules"))
		# The HR step that needs nothing but ERPNext is still asked.
		self.assertIsNotNone(self.task(self.board(), "Set up your organisation"))
