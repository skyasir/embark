"""Read a customer's sheet, match its columns to an ERPNext doctype and check every row.

The rules come from the target doctype's own meta (required fields, Select
choices, Link targets, field types), so they cannot drift from ERPNext. Nothing
here writes to ERPNext: a check returns plain-language issues and clean,
ERPNext-shaped rows.

Values a customer fixes on screen never touch their file. They are kept as
overrides and applied on every read, so the original upload stays as sent.
"""

from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from difflib import SequenceMatcher, get_close_matches

import frappe
from frappe import _
from frappe.utils import cstr

from embark.conditions import applies

ERROR = "error"
WARNING = "warning"

# Stored issue lists are capped so one sheet of 20,000 bad rows cannot bloat
# the upload record. Group counts are always complete.
# ponytail: a paginated issue endpoint is the upgrade path if 2,000 is too few.
MAX_STORED_ISSUES = 2000
MAX_ROWS = 20000

# Masters that a fresh ERPNext site already ships (units, countries, the
# standard customer groups...). A value that exists here is accepted even when
# the customer did not upload it. Everything else, such as Warehouse or Item,
# must come from the customer's own sheets: records on the onboarding site
# belong to other companies and prove nothing about theirs.
REFERENCE_DOCTYPES = {
	"UOM",
	"Country",
	"Currency",
	"Customer Group",
	"Supplier Group",
	"Territory",
	"Item Group",
	"Warehouse Type",
	"Gender",
	"Salutation",
}

# Title of each issue group on the Check screen, and whether it blocks.
# Translated when shown, not here: this runs at import, outside any request.
ISSUE_TYPES = {
	"COLUMN_MISSING": (ERROR, "Columns we couldn't find"),
	"NO_ROWS": (ERROR, "No data found"),
	"TOO_MANY_ROWS": (ERROR, "File is too large"),
	"REQUIRED_MISSING": (ERROR, "Missing values"),
	"DUPLICATE": (ERROR, "Listed more than once"),
	"NOT_FOUND": (ERROR, "Not found"),
	"BAD_CHOICE": (ERROR, "Choices not recognised"),
	"BAD_NUMBER": (ERROR, "Not a number"),
	"BAD_WHOLE_NUMBER": (ERROR, "Must be a whole number"),
	"BAD_DATE": (ERROR, "Dates we couldn't read"),
	"BAD_EMAIL": (ERROR, "Email addresses to fix"),
	"BAD_YESNO": (ERROR, "Yes / No values to fix"),
	"TOO_LONG": (ERROR, "Too long"),
	"BAD_PHONE": (WARNING, "Phone numbers to check"),
}

YES = {"1", "y", "yes", "true", "t", "x"}
NO = {"0", "n", "no", "false", "f"}

# Day-first before month-first: the markets this is sold into write 03/04/2026
# for the 3rd of April. ISO is tried first because it is unambiguous.
DATE_FORMATS = (
	"%Y-%m-%d",
	"%d-%m-%Y",
	"%d/%m/%Y",
	"%d.%m.%Y",
	"%d-%b-%Y",
	"%d %b %Y",
	"%d-%B-%Y",
	"%d %B %Y",
	"%m/%d/%Y",
	"%Y/%m/%d",
)

# How people write ERPNext's standard units. Offered as the first suggestion,
# never applied silently: some businesses really do keep "Pcs" apart from "Nos".
UNIT_SPELLINGS = {
	"kilo": "Kg",
	"kilos": "Kg",
	"kgs": "Kg",
	"kilogram": "Kg",
	"kilograms": "Kg",
	"gm": "Gram",
	"gms": "Gram",
	"grams": "Gram",
	"grm": "Gram",
	"pcs": "Nos",
	"pc": "Nos",
	"piece": "Nos",
	"pieces": "Nos",
	"no": "Nos",
	"number": "Nos",
	"numbers": "Nos",
	"ltr": "Litre",
	"ltrs": "Litre",
	"liter": "Litre",
	"liters": "Litre",
	"litres": "Litre",
	"mtr": "Meter",
	"mtrs": "Meter",
	"metre": "Meter",
	"metres": "Meter",
	"meters": "Meter",
	"boxes": "Box",
	"sets": "Set",
	"pairs": "Pair",
	"units": "Unit",
}

EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
PHONE_RE = re.compile(r"^[+0-9()\-\s./]{5,}$")

TEXT_TYPES = {"Small Text", "Text", "Long Text", "Text Editor", "Markdown Editor", "HTML Editor"}
FLOAT_TYPES = {"Float", "Currency", "Percent"}


@dataclass
class Column:
	fieldname: str
	label: str
	fieldtype: str  # Data | Email | Phone | Text | Int | Float | Check | Date | Select | Link
	required: bool = False
	choices: list[str] = field(default_factory=list)
	link_doctype: str | None = None
	default: str | None = None
	hint: str = ""
	example: str = ""
	aliases: list[str] = field(default_factory=list)
	max_length: int = 0
	is_extra: bool = False

	def as_dict(self) -> dict:
		return {
			"fieldname": self.fieldname,
			"label": self.label,
			"fieldtype": self.fieldtype,
			"required": self.required,
			"choices": self.choices,
			"link_doctype": self.link_doctype,
			"default": self.default,
			"hint": self.hint,
			"example": self.example,
			"is_extra": self.is_extra,
		}


@dataclass
class CheckResult:
	rows: list[dict]
	issues: list[dict]
	groups: list[dict]
	keys: list[str]
	errors: int
	warnings: int


# ── Columns ───────────────────────────────────────────────────────────────────


def asked_keys(answers: dict[str, str]) -> set[str]:
	"""The questions the customer is actually being asked, given the answers.

	Conditions on anything outside this set read false: the interview has already
	ruled that branch out, so there is nothing to keep open for.
	"""
	return {
		q.name
		for q in frappe.get_all("Embark Question", filters={"is_active": 1}, fields=["name", "applies_when"])
		if applies(q.applies_when, answers)
	}


def build_columns(area, answers: dict[str, str] | None = None) -> list[Column]:
	"""Typed columns for a data step, from the target doctype's meta.

	Columns the customer's answers rule out (serial numbers for a business that
	does not track them) are left out of the template and the checks.
	"""
	meta = frappe.get_meta(area.target_doctype)
	asked = asked_keys(answers) if answers else None
	columns = []
	for row in area.fields:
		if not applies(row.applies_when, answers or {}, asked):
			continue
		df = meta.get_field(row.fieldname)
		if df:
			fieldtype, options = _normalise_type(df.fieldtype, df.options)
			label = row.label or _(df.label)
			# A field ERPNext insists on is required here too, unless something
			# will fill it: its own default or the area's.
			required = bool(row.required or (df.reqd and not df.default and not row.default_value))
			max_length = (df.length or 140) if fieldtype in ("Data", "Email", "Phone") else 0
		else:
			fieldtype, options = _normalise_type(row.fieldtype or "Data", row.options)
			label = row.label or row.fieldname.replace("_", " ").title()
			required = bool(row.required)
			max_length = 140 if fieldtype in ("Data", "Email", "Phone") else 0

		columns.append(
			Column(
				fieldname=row.fieldname,
				label=label,
				fieldtype=fieldtype,
				required=required,
				choices=[o.strip() for o in (options or "").split("\n") if o.strip()]
				if fieldtype == "Select"
				else [],
				link_doctype=(options or "").strip() if fieldtype == "Link" else None,
				default=row.default_value or None,
				hint=row.hint or "",
				example=row.example or "",
				aliases=[a.strip() for a in (row.aliases or "").split(",") if a.strip()],
				max_length=max_length,
				is_extra=not df,
			)
		)
	return columns


def _normalise_type(fieldtype: str, options: str | None) -> tuple[str, str | None]:
	if fieldtype == "Data" and options in ("Email", "Phone"):
		return options, None
	if fieldtype in TEXT_TYPES:
		return "Text", None
	if fieldtype in FLOAT_TYPES:
		return "Float", None
	if fieldtype == "Datetime":
		return "Date", None
	if fieldtype in ("Data", "Email", "Phone", "Int", "Check", "Date", "Select", "Link"):
		return fieldtype, options
	return "Data", None


# ── Reading the file ──────────────────────────────────────────────────────────


def read_file(file_doc) -> tuple[list[str], list[tuple[int, list]]]:
	"""Return the header row and the data rows, each with its spreadsheet row number."""
	name = (file_doc.file_name or file_doc.file_url or "").lower()
	# File.get_content() tries to decode everything as text, which can quietly
	# turn an .xlsx into garbage; read the bytes ourselves.
	with open(file_doc.get_full_path(), "rb") as fh:
		content = fh.read()

	if name.endswith(".csv"):
		from frappe.utils.csvutils import read_csv_content

		table = read_csv_content(content)
	elif name.endswith(".xlsx"):
		table = _read_xlsx(content)
	elif name.endswith(".xls"):
		from frappe.utils.xlsxutils import read_xls_file_from_attached_file

		table = read_xls_file_from_attached_file(content)
	else:
		frappe.throw(_("Please upload an Excel (.xlsx or .xls) or CSV file."))

	return split_header(table)


def _read_xlsx(content: bytes) -> list[list]:
	from openpyxl import load_workbook

	wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
	try:
		return [list(r) for r in wb.worksheets[0].iter_rows(values_only=True)]
	finally:
		wb.close()


def split_header(table: list[list]) -> tuple[list[str], list[tuple[int, list]]]:
	"""Find the header row and return (headers, [(sheet_row_number, cells), ...]).

	People often put a title such as "Customer List" above the real headings,
	so the header is the first of the top rows that is at least as wide as two
	cells (or one, for a single-column sheet).
	"""
	widths = [sum(1 for c in r if not _blank(c)) for r in table[:15]]
	if not widths or max(widths) == 0:
		return [], []
	need = min(2, max(widths))
	head = next(i for i, w in enumerate(widths) if w >= need)

	headers, seen = [], {}
	for i, cell in enumerate(table[head]):
		h = cstr(cell).strip() or f"Column {i + 1}"
		seen[h] = seen.get(h, 0) + 1
		headers.append(h if seen[h] == 1 else f"{h} ({seen[h]})")

	rows = [
		(i + 1, list(r) + [None] * (len(headers) - len(r)))
		for i, r in enumerate(table[head + 1 :], start=head + 1)
		if any(not _blank(c) for c in r)
	]
	return headers, rows


def _blank(value) -> bool:
	return value is None or (isinstance(value, str) and not value.strip())


# ── Matching columns ──────────────────────────────────────────────────────────


def normalise(text: str) -> str:
	return re.sub(r"[^a-z0-9]", "", cstr(text).lower())


def match_columns(headers: list[str], columns: list[Column]) -> dict[str, str]:
	"""Suggest which of the customer's headings feeds which field.

	Exact matches on label, fieldname or a configured alias come first, then a
	close fuzzy match. Each field takes at most one heading.
	"""
	names = {
		c.fieldname: {normalise(c.label), normalise(c.fieldname), *map(normalise, c.aliases)} for c in columns
	}
	result: dict[str, str] = {}
	used: set[str] = set()

	for h in headers:
		n = normalise(h)
		for c in columns:
			if n and c.fieldname not in used and n in names[c.fieldname]:
				result[h] = c.fieldname
				used.add(c.fieldname)
				break

	for h in headers:
		n = normalise(h)
		if h in result or len(n) < 3:
			continue
		score, fieldname = max(
			(
				(SequenceMatcher(None, n, name).ratio(), c.fieldname)
				for c in columns
				if c.fieldname not in used
				for name in names[c.fieldname]
			),
			default=(0.0, None),
		)
		if fieldname and score >= 0.82:
			result[h] = fieldname
			used.add(fieldname)

	return result


# ── Checking rows ─────────────────────────────────────────────────────────────


def check(
	columns: list[Column],
	headers: list[str],
	rows: list[tuple[int, list]],
	column_map: dict[str, str],
	overrides: dict,
	known: dict[str, dict[str, str]],
	key_field: str,
	target_doctype: str,
	link_areas: dict[str, str] | None = None,
) -> CheckResult:
	"""Coerce every mapped cell, then check required, duplicate and link rules.

	``known`` maps a link doctype to {lowercased value: canonical value} for the
	values the customer's other sheets (and the reference masters) provide.
	``link_areas`` names the step that supplies each link doctype, for hints.
	"""
	issues: list[dict] = []
	by_field = {c.fieldname: c for c in columns}
	source = {
		fieldname: headers.index(h)
		for h, fieldname in column_map.items()
		if h in headers and fieldname in by_field
	}
	replace = overrides.get("*") or {}

	for c in columns:
		if c.required and c.fieldname not in source and not c.default:
			issues.append(
				_issue(
					"COLUMN_MISSING",
					c,
					None,
					"",
					_("We couldn't find a column for {0}.").format(c.label),
					_(
						"Pick the matching column in your file below, or add a {0} column and upload again."
					).format(c.label),
				)
			)

	if not rows:
		issues.append(
			_issue(
				"NO_ROWS",
				None,
				None,
				"",
				_("Your file has headings but no rows of data."),
				_("Fill in at least one row below the headings."),
			)
		)
	if len(rows) > MAX_ROWS:
		issues.append(
			_issue(
				"TOO_MANY_ROWS",
				None,
				None,
				"",
				_("Your file has {0} rows; the limit is {1}.").format(len(rows), MAX_ROWS),
				_("Split the file, or ask your consultant to import it directly."),
			)
		)
		rows = rows[:MAX_ROWS]

	# Pass 1: read and coerce every cell.
	clean_rows = []
	for sheet_row, cells in rows:
		fixes = overrides.get(str(sheet_row)) or {}
		out = {"_row": sheet_row}
		for c in columns:
			if c.fieldname in fixes:
				raw = fixes[c.fieldname]
			elif c.fieldname in source:
				raw = cells[source[c.fieldname]]
				swap = (replace.get(c.fieldname) or {}).get(cstr(raw).strip())
				if swap is not None:
					raw = swap
			else:
				raw = None
			value, code, detail = coerce(c, raw)
			out[c.fieldname] = value
			if code:
				issues.append(_issue(code, c, sheet_row, _display(raw), *_message(code, c, raw, detail)))
		clean_rows.append(out)

	# Duplicates on the unique field, where the sheet has one. Some do not: an
	# item may appear on several price lists.
	keys, first_seen = [], {}
	for r in clean_rows if key_field else []:
		key = cstr(r.get(key_field)).strip()
		if not key:
			continue
		folded = key.lower()
		if folded in first_seen:
			c = by_field.get(key_field)
			issues.append(
				_issue(
					"DUPLICATE",
					c,
					r["_row"],
					key,
					_("'{0}' appears more than once.").format(key),
					_("Each record should be listed once. It first appears on row {0}.").format(
						first_seen[folded]
					),
				)
			)
		else:
			first_seen[folded] = r["_row"]
			keys.append(key)

	# Pass 2: links, now that this sheet's own keys are known (an Item Group can
	# name another row of the same sheet as its parent).
	own = {k.lower(): k for k in keys}
	for c in columns:
		if c.fieldtype != "Link" or c.link_doctype == "Company":
			continue
		pool = dict(known.get(c.link_doctype) or {})
		if c.link_doctype == target_doctype:
			pool.update(own)
		for r in clean_rows:
			value = r.get(c.fieldname)
			if value in (None, ""):
				continue
			canonical = pool.get(cstr(value).lower())
			if canonical is not None:
				r[c.fieldname] = canonical
				continue
			close = _suggest(c.link_doctype, cstr(value), pool)
			issues.append(
				_issue(
					"NOT_FOUND",
					c,
					r["_row"],
					cstr(value),
					_("'{0}' is not a known {1}.").format(value, _(c.link_doctype)),
					_not_found_hint(c.link_doctype, close, pool, (link_areas or {}).get(c.link_doctype)),
					suggestions=close,
				)
			)

	errors = sum(1 for i in issues if i["severity"] == ERROR)
	return CheckResult(
		rows=clean_rows,
		issues=issues,
		groups=group_issues(issues),
		keys=keys,
		errors=errors,
		warnings=len(issues) - errors,
	)


def _suggest(doctype: str, value: str, pool: dict[str, str]) -> list[str]:
	if doctype == "UOM":
		unit = pool.get((UNIT_SPELLINGS.get(value.lower().rstrip(".")) or "").lower())
		if unit:
			# A known spelling is the answer; fuzzy look-alikes (Kilowatt for Kilo) are noise.
			return [unit]
	return get_close_matches(value, list(pool.values()), n=3, cutoff=0.6)


def _not_found_hint(doctype: str, close: list[str], pool: dict[str, str], area: str | None) -> str:
	if close:
		return _("Did you mean {0}?").format(", ".join(f"'{x}'" for x in close))
	if area:
		return _("Add it to your {0} step, or use one that's already there.").format(_(area))
	# Tree roots ("All Customer Groups") are valid parents but poor examples.
	sample = sorted(v for v in pool.values() if not v.startswith("All "))[:6]
	if sample:
		return _("Use one of ERPNext's standard values, such as {0}.").format(", ".join(sample))
	return _("Leave it empty or ask your consultant.")


def coerce(c: Column, raw) -> tuple[object, str | None, str]:
	"""Return (clean value, issue code or None, detail for the message)."""
	if _blank(raw):
		if c.default in (None, ""):
			return (None, "REQUIRED_MISSING", "") if c.required else (None, None, "")
		raw = c.default

	t = c.fieldtype
	if t == "Check":
		if isinstance(raw, bool):
			return int(raw), None, ""
		s = _number_text(raw).lower()
		if s in YES:
			return 1, None, ""
		if s in NO:
			return 0, None, ""
		return raw, "BAD_YESNO", ""

	if t in ("Int", "Float"):
		number = _to_number(raw)
		if number is None:
			return raw, "BAD_NUMBER", ""
		if t == "Int":
			if number != int(number):
				return raw, "BAD_WHOLE_NUMBER", ""
			return int(number), None, ""
		return number, None, ""

	if t == "Date":
		if isinstance(raw, datetime):
			return raw.date().isoformat(), None, ""
		if isinstance(raw, date):
			return raw.isoformat(), None, ""
		s = cstr(raw).strip()
		for fmt in DATE_FORMATS:
			try:
				return datetime.strptime(s, fmt).date().isoformat(), None, ""
			except ValueError:
				continue
		return raw, "BAD_DATE", ""

	s = _number_text(raw) if isinstance(raw, (int, float)) else cstr(raw).strip()
	if t == "Text":
		return s, None, ""
	if t == "Select":
		for choice in c.choices:
			if choice.lower() == s.lower():
				return choice, None, ""
		return s, "BAD_CHOICE", ""
	if c.max_length and len(s) > c.max_length:
		return s, "TOO_LONG", str(c.max_length)
	if t == "Email" and not EMAIL_RE.match(s):
		return s, "BAD_EMAIL", ""
	if t == "Phone" and not PHONE_RE.match(s):
		return s, "BAD_PHONE", ""
	return s, None, ""


def _number_text(raw) -> str:
	"""Excel hands back 9876543210 as 9876543210.0; show it the way it was typed."""
	if isinstance(raw, float) and raw.is_integer():
		return str(int(raw))
	return cstr(raw).strip()


def _to_number(raw) -> float | None:
	if isinstance(raw, bool):
		return None
	if isinstance(raw, (int, float)):
		return float(raw)
	s = re.sub(r"[,\s]", "", cstr(raw))
	s = re.sub(r"^[^\d\-.]+", "", s)  # currency symbols in front: ₹1200, $15
	try:
		return float(s)
	except ValueError:
		return None


def _message(code: str, c: Column, raw, detail: str) -> tuple[str, str]:
	label = c.label
	if code == "REQUIRED_MISSING":
		return _("{0} is empty.").format(label), _("Every row needs a {0}.").format(label)
	if code == "BAD_CHOICE":
		return (
			_("'{0}' isn't one of the choices for {1}.").format(_display(raw), label),
			_("Use one of: {0}.").format(", ".join(c.choices[:12])),
		)
	if code == "BAD_NUMBER":
		return _("'{0}' is not a number.").format(_display(raw)), _("Enter digits only, for example 1250.50.")
	if code == "BAD_WHOLE_NUMBER":
		return _("'{0}' must be a whole number.").format(_display(raw)), _("Remove the decimals.")
	if code == "BAD_DATE":
		return _("We couldn't read the date '{0}'.").format(_display(raw)), _("Write dates as 31-12-2026.")
	if code == "BAD_EMAIL":
		return _("'{0}' is not a valid email address.").format(_display(raw)), _(
			"It should look like name@company.com."
		)
	if code == "BAD_YESNO":
		return _("'{0}' should be Yes or No.").format(_display(raw)), _("Use Yes or No.")
	if code == "TOO_LONG":
		return (
			_("{0} is longer than {1} characters.").format(label, detail),
			_("Shorten it; the full text can go in the description."),
		)
	if code == "BAD_PHONE":
		return _("'{0}' doesn't look like a phone number.").format(_display(raw)), _(
			"Use digits, spaces, + and - only."
		)
	return code, ""


def _display(raw) -> str:
	if isinstance(raw, (datetime, date)):
		return raw.strftime("%d-%m-%Y")
	return _number_text(raw) if isinstance(raw, (int, float)) else cstr(raw).strip()


def _issue(code, column: Column | None, row, value, message, hint, suggestions=None) -> dict:
	issue = {
		"code": code,
		"severity": ISSUE_TYPES[code][0],
		"row": row,
		"fieldname": column.fieldname if column else None,
		"label": column.label if column else None,
		"value": value,
		"message": message,
		"hint": hint,
	}
	if suggestions:
		issue["suggestions"] = suggestions
	return issue


def group_issues(issues: list[dict]) -> list[dict]:
	"""One card per (problem, column) on the Check screen, errors first."""
	groups: dict[tuple, dict] = {}
	for i in issues:
		key = (i["code"], i["fieldname"])
		g = groups.get(key)
		if not g:
			g = groups[key] = {
				"code": i["code"],
				"severity": i["severity"],
				"fieldname": i["fieldname"],
				"label": i["label"],
				"title": _(ISSUE_TYPES[i["code"]][1]),
				"hint": i["hint"],
				"count": 0,
				"values": {},
			}
		g["count"] += 1
		if i["row"] is not None and i["value"] not in (None, ""):
			g["values"][i["value"]] = g["values"].get(i["value"], 0) + 1

	out = []
	for g in groups.values():
		# The distinct wrong values, most frequent first: fixing one of these
		# fixes every row that uses it.
		g["values"] = [
			{"value": v, "count": n} for v, n in sorted(g["values"].items(), key=lambda x: -x[1])[:50]
		]
		out.append(g)
	return sorted(out, key=lambda g: (g["severity"] != ERROR, -g["count"]))
